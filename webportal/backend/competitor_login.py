"""Server-side Playwright automation for the 8 tracked competitor
smallcases -- mirrors smallcase_login.py's mechanics exactly (own browser
profile/session/lock, same cwd-dependent proxy pattern, see that module's
docstring for the full story), against the SEPARATE competitor account
(holds real subscriptions to these 8).

Two other approaches were tried first and both hit hard walls smallcase
itself puts up, confirmed live (2026-10-01):
  - Replaying smallcase's own internal JSON APIs (/sam/subscriptions/v2 etc.)
    from a browser-bookmarklet 401s with "Invalid csrf token provided" even
    after replicating every header the real page sends -- an invisible
    anti-bot check only the real page's own JS can satisfy.
  - Loading each competitor's page in a hidden iframe (same bookmarklet) is
    blocked outright by smallcase's own CSP: "frame-ancestors *.juspay.in".
Real server-side Playwright navigation sidesteps both: it's a genuine
top-level page load, not an API replay or an iframe, so the real page
authenticates itself exactly as it would for a human visiting normally. The
parsing logic below is a direct port of the bookmarklet's (JS) version,
confirmed against real rendered output from a subscribed account."""
import asyncio
import os
import re
from pathlib import Path

import httpx
from _shared_http import SHARED_SSL_CONTEXT
from playwright.async_api import async_playwright

_PROFILE_DIR = Path(__file__).parent / "smallcase_session_profile_competitor"
_LOGIN_URL = "https://www.smallcase.com/smallcase/x-SCMO_0029/constituents"
_PROXY_BASE = "http://127.0.0.1:8001/api/admin"


def _should_proxy() -> bool:
    return Path.cwd().resolve() != _PROFILE_DIR.parent.resolve()


async def _proxy(method: str, path: str, auth_header: str | None = None, **kwargs) -> dict:
    headers = {"Authorization": auth_header} if auth_header else {}
    async with httpx.AsyncClient(
            verify=SHARED_SSL_CONTEXT,timeout=kwargs.pop("timeout", 180)) as client:
        resp = await client.request(method, f"{_PROXY_BASE}{path}", headers=headers, **kwargs)
        return resp.json()

# Own, independent session state -- deliberately NOT shared with
# smallcase_login.py's globals, so the two accounts never contend for the
# same lock/page/profile.
_lock = asyncio.Lock()
_pw = None
_browser_context = None
_page = None
_login_modal = None


async def _close_browser_locked():
    global _pw, _browser_context, _page, _login_modal
    try:
        if _browser_context is not None:
            await _browser_context.close()
    except Exception as e:
        print(f"[competitor_login] close_browser: {e}")
    try:
        if _pw is not None:
            await _pw.stop()
    except Exception as e:
        print(f"[competitor_login] close_browser (pw.stop): {e}")
    _pw = None
    _browser_context = None
    _page = None
    _login_modal = None


async def close_browser(auth_header: str | None = None) -> dict:
    if _should_proxy():
        return await _proxy("POST", "/competitor-login/close-browser", auth_header=auth_header)
    async with _lock:
        await _close_browser_locked()
    return {"ok": True, "message": "Browser released."}


async def _ensure_page():
    global _pw, _browser_context, _page
    if _page is not None:
        return _page
    _pw = await async_playwright().start()
    _browser_context = await _pw.chromium.launch_persistent_context(
        str(_PROFILE_DIR), headless=True, viewport={"width": 1400, "height": 1000},
    )
    _page = _browser_context.pages[0] if _browser_context.pages else await _browser_context.new_page()
    return _page


async def _debug_dump(tag: str):
    try:
        page = await _ensure_page()
        await page.screenshot(path=str(Path(__file__).parent / f"competitor_debug_{tag}.png"))
        content = await page.content()
        (Path(__file__).parent / f"competitor_debug_{tag}.html").write_text(content, encoding="utf-8")
    except Exception as e:
        print(f"[competitor_login] debug dump failed: {e}")


async def _is_logged_in() -> bool:
    try:
        page = await _ensure_page()
        await page.goto(_LOGIN_URL, wait_until="load", timeout=20000)
        await page.wait_for_timeout(1500)
        body = await page.inner_text("body")
        return "Subscribed" in body
    except Exception as e:
        print(f"[competitor_login] login check exception: {e}")
        return False


async def _start_login_locked(phone: str) -> dict:
    global _login_modal
    page = await _ensure_page()
    await page.goto(_LOGIN_URL, wait_until="load", timeout=30000)
    await page.wait_for_timeout(1500)

    body = await page.inner_text("body")
    if "Subscribed" in body:
        return {"ok": True, "already_logged_in": True, "message": "Already logged in to smallcase (competitor account)."}

    login_btn = page.locator("button:has-text('Login')").first
    if await login_btn.count() == 0:
        return {"ok": False, "error": "Could not find a Login button on the smallcase page."}
    await login_btn.click(timeout=5000)
    await page.wait_for_timeout(1000)

    phone_input = page.locator("input[data-testid='test-login-phone-number-input']").first
    if await phone_input.count() == 0:
        phone_input = page.locator("input[placeholder='Your phone number']").first
    if await phone_input.count() == 0:
        await _debug_dump("phone_input_not_found")
        return {"ok": False, "error": "Could not find the phone number field."}
    await phone_input.fill(phone)

    ancestor = phone_input.locator("xpath=ancestor::*[contains(@class, 'Modal-module')]").first
    class_attr = await ancestor.get_attribute("class") or ""
    modal_token = next((c for c in class_attr.split() if "Modal-module" in c), None)
    if modal_token:
        _login_modal = page.locator(f".{modal_token}").first
    else:
        _login_modal = ancestor

    get_otp_btn = page.locator("button:has-text('Get OTP')").first
    if await get_otp_btn.count() == 0:
        return {"ok": False, "error": "Could not find the 'Get OTP' button."}
    await get_otp_btn.click(timeout=5000)
    await page.wait_for_timeout(2000)

    return {"ok": True, "message": f"OTP requested for {phone}."}


async def _verify_otp_locked(otp: str) -> dict:
    page = await _ensure_page()

    if _login_modal is None:
        return {"ok": False, "error": "No login flow in progress -- click 'Get OTP' first."}
    modal = _login_modal
    if await modal.count() == 0:
        await _debug_dump("modal_not_found")
        return {"ok": False, "error": "Could not find the login modal -- it may have already closed."}

    otp_input = modal.locator("input[data-testid='otp-input-0']").first
    if await otp_input.count() == 0:
        await _debug_dump("otp_input_not_found")
        return {"ok": False, "error": "Could not find the OTP input field -- selector needs updating."}
    await otp_input.fill(otp)

    for _ in range(6):
        await page.wait_for_timeout(1000)
        if await _is_logged_in():
            return {"ok": True, "message": "Logged in to smallcase (competitor account) successfully."}

    await _debug_dump("login_not_confirmed")
    return {"ok": False, "error": "Login did not complete -- OTP may be wrong or the flow changed."}


async def start_login(phone: str, auth_header: str | None = None) -> dict:
    if _should_proxy():
        return await _proxy("POST", "/competitor-login/start", auth_header=auth_header, json={"phone": phone})
    async with _lock:
        return await _start_login_locked(phone)


async def verify_otp(otp: str, auth_header: str | None = None) -> dict:
    if _should_proxy():
        return await _proxy("POST", "/competitor-login/verify", auth_header=auth_header, json={"otp": otp})
    async with _lock:
        return await _verify_otp_locked(otp)


async def login_status(auth_header: str | None = None) -> dict:
    if _should_proxy():
        result = await _proxy("GET", "/competitor-login/status", auth_header=auth_header)
        return result.get("logged_in", False)
    async with _lock:
        return await _is_logged_in()


# ── Competitor smallcase map ───────────────────────────────────────────────────
# scid extracted from each smallcase's own public URL, given directly by the
# user. Key here is our own short internal identifier, not smallcase's.
COMPETITOR_SMALLCASE_MAP = {
    "gem_q_model":               {"scid": "SCMO_0029",   "label": "GEM-Q Model",                                "manager": "Windmill Capital"},
    "consumer_durables_stars":   {"scid": "QURETR_0009", "label": "Consumer Durables Stars Tracker",            "manager": "Quantace Research"},
    "omni_ai_tech":              {"scid": "OMNNM_0014",  "label": "Omni AI-Tech Global-Artificial Intelligence Theme", "manager": "Omniscience Capital"},
    "ai_data_center":            {"scid": "GRINMX_0003", "label": "AI & Data Center Theme",                     "manager": "Growth Investing"},
    "wright_innovation":         {"scid": "WRTNM_0004",  "label": "Wright Innovation Theme",                    "manager": "Wright Research"},
    "nirivantes_techwave":       {"scid": "NVNTNM_0001", "label": "Nirivantes TechWave Select Theme",           "manager": "NIRIVANTES RESEARCH"},
    "caprize_earnings_momentum": {"scid": "CAPIMO_0001", "label": "Caprize Earnings Momentum Portfolio Fundamental", "manager": "Caprize Investment Managers"},
    "caprize_midcap_smallcap":   {"scid": "CAPINM_0001", "label": "Midcap and Smallcap - Portfolio Size Model",  "manager": "Caprize Investment Managers"},
}


# ── Page-text parsing (direct port of dashboard-bookmarklet-source.js) ────────
# Real layout confirmed live (2026-10-01, subscribed account): the
# /constituents page renders CAGR, "Rebalance timeline", and the full
# "Stocks & Weights" table all on ONE page (no tab-click needed). Stock rows
# are grouped under sector sub-headings ("Biotechnology", "14.28", "Sai Life
# Sciences Limited", "7.14", "Acutaas Chemicals Ltd", "7.14", ...) -- sector
# name + sector total weight, then each real company + its own weight. Real
# NSE company names reliably end in "Limited"/"Ltd" (confirmed across all 14
# of GEM-Q Model's holdings); sector names never do, which is what
# distinguishes a sector-aggregate row from an actual stock row without
# needing to track "how many stocks in this sector".
_STOCK_NAME_RE = re.compile(r"\b(Limited|Ltd\.?)$", re.IGNORECASE)
_NUMERIC_RE = re.compile(r"^[\d.]+$")
_DATE_RE = re.compile(r"^\d{1,2} [A-Za-z]{3}, \d{4}$")


def _parse_stocks_and_weights(lines: list[str]) -> list[dict]:
    """Every (name, number) pair in this section is EITHER a sector header
    (name + that sector's aggregate weight, e.g. "Biotechnology", "14.28")
    or an actual stock (name + its own weight, e.g. "Sai Life Sciences
    Limited", "7.14") -- distinguished by the same Ltd/Limited suffix check.
    Each stock is tagged with whichever sector header most recently
    preceded it, giving real per-stock sector data for free from text
    that's already being read for weights."""
    stocks = []
    current_sector = None
    i = 0
    while i < len(lines) - 1:
        if _NUMERIC_RE.match(lines[i + 1]):
            if _STOCK_NAME_RE.search(lines[i]):
                stocks.append({"name": lines[i], "weight": float(lines[i + 1]), "sector": current_sector})
            else:
                current_sector = lines[i]
            i += 2
        else:
            i += 1
    return stocks


_CAP_LABELS = ("Largecap", "Midcap", "Smallcap")


def _parse_market_cap_mix(lines: list[str]) -> dict:
    if "Holdings Distribution" not in lines:
        return {}
    start = lines.index("Holdings Distribution")
    end = len(lines)
    for marker in ("Stocks & Weights", "Stocks & Segments"):
        if marker in lines[start:]:
            end = min(end, lines.index(marker, start))
    result = {}
    i = start + 1
    while i < end - 1:
        if lines[i] in _CAP_LABELS and re.match(r"^[\d.]+%$", lines[i + 1]):
            result[lines[i]] = float(lines[i + 1].rstrip("%"))
            i += 2
        else:
            i += 1
    return result


_RANGE_RE = re.compile(r"^(\d{1,2} [A-Za-z]{3}, \d{4}) - (\d{1,2} [A-Za-z]{3}, \d{4})$")
_COUNT_RE = re.compile(r"^\+(\d+) Rebalances?$")


def _parse_rebalance_timeline(lines: list[str], start_idx: int, end_idx: int) -> list[dict]:
    """smallcase's own timeline widget only ever renders per-stock-change
    detail for the last ~3 rebalances -- everything older is collapsed into
    one "+N Rebalances" summary spanning a date range, confirmed live
    (2026-10-01) that this collapsed bucket is NOT clickable/expandable (no
    DOM change after clicking it), so individual historical rebalance dates
    in that range genuinely aren't recoverable from this page. Captured here
    as a single "summary" event instead of being silently dropped, alongside
    the "smallcase went Live" launch marker, so the full since-inception
    picture (coarse as it is) is visible rather than looking truncated."""
    events = []
    for i in range(start_idx, end_idx - 1):
        if _DATE_RE.match(lines[i]) and lines[i + 1] == "Constituents updated":
            added = removed = None
            if i + 2 < len(lines) and re.match(r"^\+\d+$", lines[i + 2]):
                added = int(lines[i + 2][1:])
            if i + 3 < len(lines) and re.match(r"^-\d+$", lines[i + 3]):
                removed = int(lines[i + 3][1:])
            events.append({"date": lines[i], "assetsAdded": added, "assetsRemoved": removed, "type": "rebalance"})
        elif _DATE_RE.match(lines[i]) and i + 1 < len(lines) and "went Live" in lines[i + 1]:
            events.append({"date": lines[i], "type": "launch"})
        elif i + 1 < len(lines):
            m_range = _RANGE_RE.match(lines[i])
            m_count = _COUNT_RE.match(lines[i + 1]) if m_range else None
            if m_range and m_count:
                events.append({
                    "date": m_range.group(2), "type": "summary",
                    "rangeStart": m_range.group(1), "rangeEnd": m_range.group(2),
                    "rebalanceCount": int(m_count.group(1)),
                })
    return events


def _parse_competitor_page(body_text: str) -> dict:
    lines = [l.strip() for l in body_text.split("\n") if l.strip()]
    result: dict = {}

    if "CAGR" in lines:
        cagr_idx = lines.index("CAGR")
        if cagr_idx + 1 < len(lines):
            m = re.search(r"([\d.]+)%", lines[cagr_idx + 1])
            if m:
                result["cagr"] = float(m.group(1))

    if "Rebalance timeline" in lines:
        start = lines.index("Rebalance timeline")
        end = len(lines)
        for marker in ("Portfolio Report", "Get portfolio report"):
            if marker in lines[start:]:
                end = lines.index(marker, start)
                break
        result["rebalanceTimeline"] = _parse_rebalance_timeline(lines, start, end)

    if "Weightage (%)" in lines:
        stocks_start = lines.index("Weightage (%)") + 1
        stocks_end = len(lines)
        for idx in range(stocks_start, len(lines)):
            if lines[idx].startswith("You can also edit constituents"):
                stocks_end = idx
                break
        result["stocks"] = _parse_stocks_and_weights(lines[stocks_start:stocks_end])

    result["marketCapMix"] = _parse_market_cap_mix(lines)

    return result


# ── Portfolio report PDF (richer, more reliable than page-text scraping) ─────
# Confirmed live (2026-10-01): the "Get portfolio report" PDF download
# (same download-capture technique as smallcase_login._fetch_rebalance_
# report_locked) works on competitor pages too, and uses the IDENTICAL
# report format as our own smallcase's reports -- "Portfolio Constituents"
# (full current holdings, WITH each stock's exact market-cap Holding Type,
# not just a page-level aggregate %) and "Latest Rebalance Update" (exact
# per-stock weight deltas on the most recent rebalance). Crucially, the PDF
# password is tied to the SUBSCRIBER ACCOUNT, not each report's publisher --
# confirmed live that the one password (COMPETITOR_PORTFOLIO_PDF_PASSWORD)
# decrypts Windmill Capital's GEM-Q Model report fine, even though neither
# PORTFOLIO_PDF_PASSWORD (our own account's) nor a blank password worked.
# This means the SAME parse_rebalance_update() parser already built and
# proven for our own reports (smallcase_login.py) works here UNCHANGED.
_CAGR_LINE_RE = re.compile(r"^([\d.]+)%CAGR")
_CONSTITUENT_LINE_RE = re.compile(r"^(.+?)\s+(Smallcap|Midcap|Largecap|Multicap)\s+([\d.]+)%$")


def _parse_portfolio_constituents(lines: list[str]) -> list[dict]:
    """Confirmed live (2026-10-01, Nirivantes TechWave Select Theme): pypdf's
    extract_text() sometimes wraps a longer company name onto its own line,
    separate from the "Segment Weight%" that follows -- a plain per-line
    regex silently dropped 14 of 18 real holdings for that one report. Each
    non-matching line is retried joined with the NEXT line before being
    given up on, which recovers the wrapped-name case without needing to
    touch smallcase_login.parse_rebalance_update (a separately-shared,
    already-proven parser for our own reports -- safer to leave that one
    alone than risk regressing it for this)."""
    try:
        start = next(i for i, l in enumerate(lines) if l == "Portfolio Constituents")
    except StopIteration:
        return []
    end = next((i for i in range(start, len(lines)) if "Click here to view sectorwise" in lines[i]), len(lines))
    section = lines[start + 1:end]
    out = []
    i = 0
    while i < len(section):
        m = _CONSTITUENT_LINE_RE.match(section[i])
        if m:
            name, cap_segment, weight = m.groups()
            out.append({"name": name.strip(), "capSegment": cap_segment, "weight": float(weight)})
            i += 1
            continue
        if i + 1 < len(section):
            m2 = _CONSTITUENT_LINE_RE.match(f"{section[i]} {section[i + 1]}")
            if m2:
                name, cap_segment, weight = m2.groups()
                out.append({"name": name.strip(), "capSegment": cap_segment, "weight": float(weight)})
                i += 2
                continue
        i += 1
    return out


def _extract_report_cagr(lines: list[str]) -> float | None:
    for l in lines:
        m = _CAGR_LINE_RE.match(l)
        if m:
            return float(m.group(1))
    return None


def _extract_launch_date(lines: list[str]) -> str | None:
    try:
        idx = lines.index("Launch Date")
    except ValueError:
        return None
    if idx + 1 >= len(lines):
        return None
    try:
        from datetime import datetime as _dt
        return _dt.strptime(lines[idx + 1], "%B %d, %Y").strftime("%d %b %Y")
    except Exception:
        return None


async def _fetch_competitor_report_text(page, scid: str) -> str | None:
    """Same download-capture technique as smallcase_login._fetch_rebalance_
    report_locked, reused against a competitor's own page -- returns the
    decrypted report's full text, or None if the button isn't found, the
    download doesn't arrive, or decryption fails (any of which just means
    callers fall back to page-scraped data instead, see
    _fetch_competitor_snapshot_locked)."""
    report_btn = page.locator("text=/Get portfolio report/i").first
    if await report_btn.count() == 0:
        return None
    await report_btn.click()
    await page.wait_for_timeout(1200)

    view_btn = page.locator("text=/View report/i").first
    if await view_btn.count() == 0:
        return None

    download_holder: dict = {}

    async def _save_download(d):
        download_holder["download"] = d

    def _on_new_page(p):
        p.on("download", lambda d: asyncio.create_task(_save_download(d)))

    page.on("download", lambda d: asyncio.create_task(_save_download(d)))
    page.context.on("page", _on_new_page)

    await view_btn.click()
    for _ in range(15):
        await asyncio.sleep(1)
        if "download" in download_holder:
            break
    if "download" not in download_holder:
        return None

    tmp_path = Path(__file__).parent / f"_tmp_competitor_report_{scid}.pdf"
    try:
        await download_holder["download"].save_as(str(tmp_path))
        raw = tmp_path.read_bytes()
    finally:
        try:
            tmp_path.unlink(missing_ok=True)
        except Exception:
            pass

    try:
        from pypdf import PdfReader
        import io as _io
        reader = PdfReader(_io.BytesIO(raw))
        if reader.is_encrypted:
            if reader.decrypt(os.environ.get("COMPETITOR_PORTFOLIO_PDF_PASSWORD", "")) == 0:
                print(f"[competitor_login] report PDF password did not work for {scid}")
                return None
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    except Exception as e:
        print(f"[competitor_login] report PDF parse failed for {scid}: {e}")
        return None


async def _fetch_rebalance_timeline_xlsx(page, scid: str) -> bytes | None:
    """The page-scraped "Rebalance timeline" only ever shows the last ~3
    rebalances + a collapsed "+N Rebalances" summary for everything older
    (confirmed non-expandable, see _parse_rebalance_timeline's docstring).
    BUT the constituents page also has a "Download rebalance timeline"
    button (confirmed live 2026-10-01, no password) that exports an .xlsx
    with a "Historical Constituents" sheet listing every historical
    rebalance period since inception with its full stock list and weights
    -- this is the real fix for "since inception" competitor history."""
    try:
        btn = page.locator("[aria-label='Download rebalance timeline'], [title='Download rebalance timeline']").first
        if await btn.count() == 0:
            return None
        async with page.expect_download(timeout=15000) as dl_info:
            await btn.click()
        download = await dl_info.value
        tmp_path = _PROFILE_DIR.parent / f"_tmp_timeline_{scid}.xlsx"
        await download.save_as(str(tmp_path))
        data = tmp_path.read_bytes()
        tmp_path.unlink(missing_ok=True)
        return data
    except Exception as e:
        print(f"[competitor_login] rebalance-timeline xlsx download failed for {scid}: {e}")
        return None


def _parse_historical_constituents_xlsx(xlsx_bytes: bytes) -> list[dict]:
    """Parses the "Historical Constituents" sheet (Date Range / Constituents
    / Weightage columns, one row per stock, "Date Range" only filled on the
    first row of each period's block -- forward-filled here) into
    chronologically-ascending periods: [{"date": "YYYY-MM-DD" (period start),
    "stocks": [{"name", "weight"} (weight as a 0-100 percentage, the sheet
    stores it as a 0-1 fraction)]}]."""
    import io
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    sheet = None
    for name in wb.sheetnames:
        if "historical constituent" in name.lower():
            sheet = wb[name]
            break
    if sheet is None:
        return []

    periods: list[dict] = []
    current_date = None
    current_stocks: list[dict] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 3:
            continue
        date_val, name_val, weight_val = row[0], row[1], row[2]
        if date_val:
            if current_date and current_stocks:
                periods.append({"date": current_date, "stocks": current_stocks})
            start = str(date_val).split(" to ")[0].strip()
            current_date = start[:10]
            current_stocks = []
        if name_val:
            try:
                weight_pct = round(float(weight_val or 0) * 100, 4)
            except (TypeError, ValueError):
                weight_pct = 0.0
            current_stocks.append({"name": str(name_val).strip(), "weight": weight_pct})
    if current_date and current_stocks:
        periods.append({"date": current_date, "stocks": current_stocks})
    periods.sort(key=lambda p: p["date"])
    return periods


def diff_historical_constituents(periods: list[dict]) -> list[dict]:
    """Turns the chronologically-ascending period list from
    _parse_historical_constituents_xlsx into the same new/increased/
    decreased/removed-per-date shape as our own basket-rebalance-history
    endpoint (minus nseCode, which the caller resolves and attaches, same
    as it already does for the plain `stocks` list) -- matched by stock
    NAME since names are stable across one competitor's own periods (no
    cross-portfolio ambiguity like the timing-insights nseCode match
    needs). Returned newest-first, mirroring ourRebalanceHistory's order."""
    history = []
    prior_weight: dict = {}
    for period in periods:
        changes = []
        counts = {"new": 0, "increased": 0, "decreased": 0, "removed": 0}
        current_names = {s["name"] for s in period["stocks"]}
        for s in period["stocks"]:
            name = s["name"]
            new_w = s["weight"]
            old_w = prior_weight.get(name)
            if old_w is None:
                status = "new"
            elif new_w > old_w:
                status = "increased"
            elif new_w < old_w:
                status = "decreased"
            else:
                status = None
            if status:
                counts[status] += 1
                changes.append({"name": name, "oldWeight": old_w, "newWeight": new_w, "status": status})
        for name, old_w in prior_weight.items():
            if name not in current_names:
                counts["removed"] += 1
                changes.append({"name": name, "oldWeight": old_w, "newWeight": 0, "status": "removed"})
        if changes:
            history.append({"date": period["date"], "counts": counts, "changes": changes})
        prior_weight = {s["name"]: s["weight"] for s in period["stocks"]}
    history.reverse()
    return history


async def _fetch_competitor_performance_series(page, scid: str) -> list[dict]:
    """Real indexed performance time series for this competitor -- same
    /sam/graph/performance endpoint our own baskets' daily-fetch already
    uses, confirmed live (2026-10-01) to work on cookies alone (no CSRF
    needed, unlike /sam/subscriptions/v2 and friends -- this endpoint was
    never the blocker). Fetched via page.request so it carries the
    logged-in browser context's cookies. Lets the frontend compute returns
    for whatever period it wants (1M/3M/6M/1Y/3Y/since-inception), using the
    exact same computeTenureReturn() logic already used for our own
    baskets -- so "CAGR" always comes with a real, matching date range
    instead of a bare unexplained percentage."""
    try:
        url = f"https://api.smallcase.com/sam/graph/performance?currency=INR&duration=max&scids[]={scid}"
        resp = await page.request.get(url)
        if resp.status != 200:
            return []
        body = await resp.json()
        points = (body.get("data", {}).get(scid) or {}).get("points", [])
        return [{"date": p["date"][:10], "value": p["price"]} for p in points if p.get("price") is not None]
    except Exception as e:
        print(f"[competitor_login] performance series fetch failed for {scid}: {e}")
        return []


async def _fetch_competitor_snapshot_locked(key: str) -> dict:
    if key not in COMPETITOR_SMALLCASE_MAP:
        return {"ok": False, "error": f"Unknown competitor: {key}"}
    cfg = COMPETITOR_SMALLCASE_MAP[key]
    scid = cfg["scid"]

    page = await _ensure_page()
    await page.goto(f"https://www.smallcase.com/smallcase/x-{scid}/constituents", wait_until="load", timeout=30000)
    await page.wait_for_timeout(2500)

    body = await page.inner_text("body")
    if "Subscribed" not in body:
        return {"ok": False, "error": "Not logged in / not subscribed to this competitor smallcase."}

    parsed = _parse_competitor_page(body)
    if not parsed.get("stocks") and parsed.get("cagr") is None and not parsed.get("rebalanceTimeline"):
        return {"ok": False, "error": "Could not find any recognizable section on this page."}

    # Prefer the portfolio-report PDF's data over the page-scrape where
    # available -- exact per-stock market-cap segment and weight, a more
    # precisely-dated CAGR, the real launch date, and (new) the exact
    # per-stock additions/removals/weight-deltas from the latest rebalance,
    # all parsed with the same proven logic our own reports already use.
    industry_sector_by_name = {s["name"]: s.get("sector") for s in (parsed.get("stocks") or [])}
    try:
        pdf_text = await _fetch_competitor_report_text(page, scid)
    except Exception as e:
        print(f"[competitor_login] report PDF fetch errored for {scid}: {e}")
        pdf_text = None

    if pdf_text:
        pdf_lines = [l.strip() for l in pdf_text.split("\n") if l.strip()]
        constituents = _parse_portfolio_constituents(pdf_lines)
        constituents_weight = sum(c["weight"] for c in constituents)
        if constituents and constituents_weight < 90:
            # Weights should sum to ~100% -- well short of that means several
            # rows still failed to parse (e.g. a name wrapped across three+
            # lines, beyond what the one-line join above recovers). Safer to
            # keep the page-scraped stock list than publish a silently
            # incomplete one as if it were authoritative.
            print(f"[competitor_login] PDF constituents for {scid} only summed to {constituents_weight:.1f}% -- keeping page-scraped stocks instead")
            constituents = []
        if constituents:
            parsed["stocks"] = [
                {"name": c["name"], "weight": c["weight"], "capSegment": c["capSegment"],
                 "sector": industry_sector_by_name.get(c["name"])}
                for c in constituents
            ]
            precise_cap_mix: dict = {}
            for c in constituents:
                precise_cap_mix[c["capSegment"]] = round(precise_cap_mix.get(c["capSegment"], 0) + c["weight"], 2)
            parsed["marketCapMix"] = precise_cap_mix
        report_cagr = _extract_report_cagr(pdf_lines)
        if report_cagr is not None:
            parsed["cagr"] = report_cagr
        parsed["launchDate"] = _extract_launch_date(pdf_lines)
        import smallcase_login
        parsed["latestRebalanceDetail"] = smallcase_login.parse_rebalance_update(pdf_text)
    else:
        print(f"[competitor_login] falling back to page-scraped stocks for {scid} (report PDF unavailable)")

    parsed["performanceSeries"] = await _fetch_competitor_performance_series(page, scid)

    try:
        xlsx_bytes = await _fetch_rebalance_timeline_xlsx(page, scid)
        parsed["historicalConstituents"] = _parse_historical_constituents_xlsx(xlsx_bytes) if xlsx_bytes else []
    except Exception as e:
        print(f"[competitor_login] historical-constituents parse failed for {scid}: {e}")
        parsed["historicalConstituents"] = []

    parsed["ok"] = True
    parsed["key"] = key
    parsed["label"] = cfg["label"]
    parsed["manager"] = cfg["manager"]
    return parsed


async def fetch_competitor_snapshot(key: str, auth_header: str | None = None) -> dict:
    if _should_proxy():
        return await _proxy("POST", f"/competitor-fetch-snapshot/{key}", auth_header=auth_header, timeout=60)
    async with _lock:
        return await _fetch_competitor_snapshot_locked(key)


async def fetch_all_competitors(auth_header: str | None = None) -> dict:
    """Admin-facing entry point: logs in-status check once, then scrapes
    every tracked competitor in one server-side run (this is what the
    Competitor Analysis page's "Fetch Competitor Data" button calls)."""
    if _should_proxy():
        return await _proxy("POST", "/competitor-fetch-all", auth_header=auth_header, timeout=300)
    async with _lock:
        if not await _is_logged_in():
            return {"ok": False, "error": "Not logged in to smallcase (competitor account) -- use the login button first."}
        results = {}
        for key in COMPETITOR_SMALLCASE_MAP:
            try:
                results[key] = await _fetch_competitor_snapshot_locked(key)
            except Exception as e:
                results[key] = {"ok": False, "error": str(e)}
        return {"ok": True, "results": results}
